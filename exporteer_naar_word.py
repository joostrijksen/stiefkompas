#!/usr/bin/env python3
"""
Stiefkompas Content Exporter
Exporteert alle MDX en TSX content naar een Word-document voor review.

Gebruik:
  pip install python-docx
  python exporteer_naar_word.py

Output: stiefkompas_content_review.docx
"""

import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── Configuratie ───────────────────────────────────────────────────────────────

SCRIPT_DIR = Path(__file__).parent
OUTPUT_FILE = SCRIPT_DIR / "stiefkompas_content_review.docx"

# MDX mappen in volgorde
MDX_SECTIONS = [
    ("Module 1", "content/v1.1/module-1"),
    ("Module 2", "content/v1.1/module-2"),
    ("Certificering", "content/v1.1/certificering"),
]

# TSX pagina's met tekst (alleen site-pagina's, geen auth/portaal)
TSX_PAGES = [
    ("Homepage", "app/(site)/page.tsx"),
    ("Methode", "app/(site)/methode/page.tsx"),
    ("Wetenschappelijke basis", "app/(site)/wetenschappelijke-basis/page.tsx"),
    ("Opleiding — Module 1", "app/(site)/opleiding/module-1/page.tsx"),
    ("Opleiding — Module 2", "app/(site)/opleiding/module-2/page.tsx"),
    ("Opleiding — Certificering", "app/(site)/opleiding/certificering/page.tsx"),
    ("Voor coaches", "app/(site)/voor-coaches/page.tsx"),
    ("Werkwijze", "app/(site)/werkwijze/page.tsx"),
    ("Inschrijven", "app/(site)/inschrijven/page.tsx"),
    ("FAQ", "app/(site)/faq/page.tsx"),
    ("Materialen & Tools (overzicht)", "app/(site)/materialen-tools/page.tsx"),
    ("Publicaties (overzicht)", "app/(site)/publicaties/page.tsx"),
    ("Publicatie — Loyaliteit & Positie", "app/(site)/publicaties/loyaliteit-positie/page.tsx"),
    ("Publicatie — Hechting & Wisselritmes", "app/(site)/publicaties/hechting-wisselritmes/page.tsx"),
    ("Publicatie — Communicatiepatronen", "app/(site)/publicaties/communicatiepatronen/page.tsx"),
    ("Publicatie — Systemisch perspectief", "app/(site)/publicaties/systemisch-perspectief/page.tsx"),
    ("Publicatie — Kompasmodel & Teams", "app/(site)/publicaties/kompasmodel-teams/page.tsx"),
]

# ── Tekst extractie ────────────────────────────────────────────────────────────

def clean_mdx(text: str) -> str:
    """Verwijder frontmatter, MDX-componenten en opmaak voor leesbare tekst."""
    # Verwijder frontmatter (--- ... ---)
    text = re.sub(r'^---[\s\S]*?---\n', '', text, flags=re.MULTILINE)
    # Verwijder JSX/HTML tags
    text = re.sub(r'<[^>]+>', '', text)
    # Verwijder import statements
    text = re.sub(r'^import .+$', '', text, flags=re.MULTILINE)
    # Verwijder code blocks (behoud inhoud)
    text = re.sub(r'```[\w]*\n([\s\S]*?)```', r'\1', text)
    # Verwijder inline code
    text = re.sub(r'`([^`]+)`', r'\1', text)
    # Verwijder afbeelding-placeholders
    text = re.sub(r'>\s*\*\*\[AFBEELDING[^\]]*\]\*\*[\s\S]*?(?=\n---|\n##|\Z)', '', text)
    # Opruimen
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def extract_strings_from_tsx(text: str) -> list[str]:
    """Haal leesbare tekst uit TSX-bestanden."""
    results = []

    # Haal strings uit JSX (tekst tussen > en <, of in quotes in arrays/objects)
    # Methode 1: tekst tussen JSX tags
    jsx_text = re.findall(r'>\s*([A-ZÀ-ÿa-z][^<>{}\n]{10,}?)\s*<', text)
    results.extend(jsx_text)

    # Methode 2: strings in objects/arrays (title/desc/definition/intro patronen)
    for field in ['title', 'desc', 'description', 'intro', 'definition', 'theory',
                  'extra', 'name', 'term', 'signals', 'elements']:
        matches = re.findall(
            rf'(?:"{field}"|{field}):\s*[`"]([^`"{{}}]{{15,}})[`"]',
            text
        )
        results.extend(matches)

    # Methode 3: template literals (backtick strings > 30 chars)
    backtick = re.findall(r'`([^`]{30,})`', text, re.DOTALL)
    for b in backtick:
        if not any(c in b for c in ['import', 'export', 'const ', 'function', '=>', '${', 'className']):
            results.append(b.strip())

    # Dedupliceer en filter
    seen = set()
    clean = []
    for r in results:
        r = r.strip()
        r = re.sub(r'\s+', ' ', r)
        if len(r) > 20 and r not in seen:
            # Filter technische strings
            if not any(x in r for x in ['px', 'className', 'style=', 'import ', 'const ', '=>', 'onClick', 'href=']):
                seen.add(r)
                clean.append(r)

    return clean


def get_mdx_title(text: str, filepath: str) -> str:
    """Haal titel uit frontmatter of bestandsnaam."""
    match = re.search(r'^title:\s*["\']?(.+?)["\']?\s*$', text, re.MULTILINE)
    if match:
        return match.group(1).strip()
    # Fallback: bestandsnaam
    return Path(filepath).stem.replace('-', ' ').title()


def parse_mdx_structure(text: str) -> list[tuple[int, str]]:
    """
    Geef lijst van (niveau, inhoud) terug voor headings en paragrafen.
    niveau: 1=H1, 2=H2, 3=H3, 0=paragraaf
    """
    clean = clean_mdx(text)
    result = []
    for line in clean.splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith('### '):
            result.append((3, line[4:].strip()))
        elif line.startswith('## '):
            result.append((2, line[3:].strip()))
        elif line.startswith('# '):
            result.append((1, line[2:].strip()))
        elif line.startswith('- ') or line.startswith('* '):
            result.append((-1, line[2:].strip()))  # bullet
        elif re.match(r'^\d+\.\s', line):
            result.append((-2, re.sub(r'^\d+\.\s', '', line).strip()))  # numbered
        elif line.startswith('>'):
            result.append((-3, line.lstrip('> ').strip()))  # blockquote
        elif line.startswith('✓') or line.startswith('✗'):
            result.append((-1, line.strip()))
        elif len(line) > 5:
            result.append((0, line.strip()))
    return result


# ── Word document opbouw ───────────────────────────────────────────────────────

def add_section_divider(doc: Document, title: str):
    """Voeg een sectie-header toe (Heading 1 stijl)."""
    p = doc.add_heading(title, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x0E, 0x2A, 0x47)


def add_subsection(doc: Document, title: str, level: int = 2):
    """Voeg een subsectie-header toe."""
    p = doc.add_heading(title, level=level)
    if p.runs:
        p.runs[0].font.color.rgb = RGBColor(0x0E, 0x2A, 0x47)


def add_body_text(doc: Document, text: str):
    """Voeg normale paragraaf toe."""
    if text.strip():
        doc.add_paragraph(text.strip())


def add_bullet(doc: Document, text: str):
    """Voeg bullet-punt toe."""
    if text.strip():
        doc.add_paragraph(text.strip(), style='List Bullet')


def add_note(doc: Document, text: str):
    """Voeg een opmerking/quote toe in italic."""
    if text.strip():
        p = doc.add_paragraph()
        run = p.add_run(text.strip())
        run.italic = True
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def add_page_break(doc: Document):
    doc.add_page_break()


def add_review_box(doc: Document):
    """Voeg lege reviewnotitie toe voor Karin."""
    p = doc.add_paragraph()
    run = p.add_run("📝 Reviewnotitie: ")
    run.bold = True
    run.font.color.rgb = RGBColor(0x17, 0xB3, 0xB0)
    p.add_run("_" * 60)


# ── MDX verwerken ──────────────────────────────────────────────────────────────

def process_mdx_file(doc: Document, filepath: Path):
    """Verwerk één MDX-bestand naar Word."""
    try:
        text = filepath.read_text(encoding='utf-8')
    except Exception as e:
        doc.add_paragraph(f"[Fout bij lezen: {e}]")
        return

    title = get_mdx_title(text, str(filepath))
    add_subsection(doc, title, level=3)

    structure = parse_mdx_structure(text)
    for level, content in structure:
        if not content:
            continue
        if level == 1:
            add_subsection(doc, content, level=4)
        elif level == 2:
            add_subsection(doc, content, level=4)
        elif level == 3:
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.bold = True
        elif level == -1:
            add_bullet(doc, content)
        elif level == -2:
            doc.add_paragraph(content, style='List Number')
        elif level == -3:
            add_note(doc, content)
        else:
            # Sla technische regels over
            if not any(x in content for x in ['className', 'import ', 'const ', '=>', 'px', '{', '}']):
                add_body_text(doc, content)

    add_review_box(doc)
    doc.add_paragraph("")  # witregel


def process_mdx_folder(doc: Document, section_name: str, folder_path: Path):
    """Verwerk alle MDX-bestanden in een map, gesorteerd op pad."""
    add_section_divider(doc, section_name)

    mdx_files = sorted(folder_path.rglob("*.mdx"))
    if not mdx_files:
        doc.add_paragraph(f"(Geen MDX-bestanden gevonden in {folder_path})")
        return

    current_chapter = None
    for mdx_file in mdx_files:
        # Hoofdstuk-groepering op basis van mapnaam
        parts = mdx_file.parts
        chapter_idx = None
        for i, part in enumerate(parts):
            if 'hoofdstuk' in part or 'afsluiting' in part or 'certificering' in part or 'module' in part:
                chapter_idx = i
        chapter = parts[chapter_idx] if chapter_idx else mdx_file.parent.name
        chapter_label = chapter.replace('-', ' ').title()

        if chapter != current_chapter:
            current_chapter = chapter
            add_subsection(doc, chapter_label, level=2)

        process_mdx_file(doc, mdx_file)

    add_page_break(doc)


# ── TSX verwerken ──────────────────────────────────────────────────────────────

def process_tsx_file(doc: Document, label: str, filepath: Path):
    """Verwerk één TSX-bestand: extraheer leesbare tekst."""
    add_subsection(doc, label, level=2)

    if not filepath.exists():
        doc.add_paragraph(f"(Bestand niet gevonden: {filepath})")
        return

    try:
        text = filepath.read_text(encoding='utf-8')
    except Exception as e:
        doc.add_paragraph(f"[Fout bij lezen: {e}]")
        return

    strings = extract_strings_from_tsx(text)

    if not strings:
        doc.add_paragraph("(Geen tekst gevonden — mogelijk dynamisch geladen)")
    else:
        for s in strings:
            # Meerdere regels = paragrafen
            for line in s.split('\n\n'):
                line = line.strip()
                if line and len(line) > 10:
                    add_body_text(doc, line)

    add_review_box(doc)
    doc.add_paragraph("")


# ── Hoofdprogramma ─────────────────────────────────────────────────────────────

def main():
    print("🚀 Stiefkompas Content Exporter")
    print(f"   Werkmap: {SCRIPT_DIR}")
    print(f"   Output:  {OUTPUT_FILE}")
    print()

    doc = Document()

    # Stijlen instellen
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    # Titelpagina
    title_p = doc.add_heading('Stiefkompas — Content Review', level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph('Exportdocument voor tekstcontrole en -aanpassing')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    doc.add_paragraph(
        'Dit document bevat alle teksten uit de Stiefkompas website en leeromgeving. '
        'Gebruik de reviewnotities (📝) om aanpassingen te markeren of commentaar toe te voegen.'
    )
    add_page_break(doc)

    # ── DEEL 1: MDX Content ──────────────────────────────────────────────────
    doc.add_heading('DEEL 1 — Lesinhoud (MDX)', level=1)
    doc.add_paragraph(
        'Dit deel bevat de volledige lesinhoud van beide modules en de certificering.'
    )
    add_page_break(doc)

    for section_name, folder_rel in MDX_SECTIONS:
        folder = SCRIPT_DIR / folder_rel
        if folder.exists():
            process_mdx_folder(doc, section_name, folder)
            print(f"  ✅ MDX: {section_name}")
        else:
            print(f"  ⚠️  Map niet gevonden: {folder}")

    # ── DEEL 2: Website Teksten (TSX) ────────────────────────────────────────
    doc.add_heading('DEEL 2 — Website teksten (TSX pagina\'s)', level=1)
    doc.add_paragraph(
        'Dit deel bevat de teksten van de publieke website-pagina\'s. '
        'De tekst is automatisch geëxtraheerd uit de TSX-bestanden.'
    )
    add_page_break(doc)

    for label, tsx_rel in TSX_PAGES:
        tsx_path = SCRIPT_DIR / tsx_rel
        process_tsx_file(doc, label, tsx_path)
        status = "✅" if tsx_path.exists() else "⚠️ "
        print(f"  {status} TSX: {label}")

    # Opslaan
    doc.save(OUTPUT_FILE)
    print()
    print(f"✅ Klaar! Bestand opgeslagen: {OUTPUT_FILE}")
    print()
    print("Volgende stap:")
    print("  1. Upload het .docx bestand naar Google Drive")
    print("  2. Klik met rechtermuisknop → 'Openen met Google Docs'")
    print("  3. Deel met Karin voor review")


if __name__ == "__main__":
    main()
